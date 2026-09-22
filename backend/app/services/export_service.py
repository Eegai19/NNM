"""Excel report generation with openpyxl."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", start_color="FF124191", end_color="FF124191")
HEADER_FONT = Font(color="FFFFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="FFD5DCE6")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _format_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if value is None:
        return ""
    if isinstance(value, (int, float, str)):
        return value
    return str(value)


def _write_sheet(sheet: Worksheet, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    sheet.append(list(headers))
    for column_index, _ in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = CELL_BORDER

    for row in rows:
        sheet.append([_format_value(value) for value in row])

    # Auto-size columns against the longest cell, capped so the sheet stays readable.
    for column_index, header in enumerate(headers, start=1):
        longest = len(str(header))
        for row in rows:
            value = row[column_index - 1]
            longest = max(longest, len(str(_format_value(value))))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(longest + 4, 12), 48)

    for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(headers)):
        for cell in row_cells:
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(sheet.max_row, 1)}"


def build_workbook(sheets: dict[str, tuple[Sequence[str], Sequence[Sequence[Any]]]]) -> bytes:
    """Render one or more sheets into an in-memory ``.xlsx`` file."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, (headers, rows) in sheets.items():
        sheet = workbook.create_sheet(title=title[:31])
        _write_sheet(sheet, headers, rows)
    if not workbook.sheetnames:  # pragma: no cover - defensive
        workbook.create_sheet("Empty")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


NODE_HEADERS = (
    "Node ID",
    "Node Name",
    "Product",
    "Circle",
    "Deployment State",
    "Overall Status",
    "Owner",
    "Lead",
    "TPM",
    "Total Activities",
    "Pending",
    "In Progress",
    "Completed",
    "Created At",
    "Updated At",
)

ACTIVITY_HEADERS = (
    "Activity ID",
    "Node",
    "Circle",
    "Product",
    "Activity",
    "Assigned To",
    "Status",
    "Start Date",
    "Completed Date",
    "Logs",
    "Remarks",
)
